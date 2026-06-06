import openpyxl

wb = openpyxl.load_workbook('output/cleaned_test_null.xlsx')
ws = wb.active

print('=== 空值处理验证 (直接读取Excel单元格) ===')
cols = {'B': '姓名', 'G': '联系电话', 'F': '填写日期', 'E': '省份'}
all_ok = True

for col_letter, col_name in cols.items():
    nan_count = 0
    empty_count = 0
    for row in range(2, ws.max_row + 1):
        cell = ws[f'{col_letter}{row}']
        val = cell.value
        if val is None or val == '':
            empty_count += 1
        elif str(val).lower() == 'nan':
            nan_count += 1
            all_ok = False
    
    status = '✅' if nan_count == 0 else '❌'
    print(f'{status} {col_name}: 空值{empty_count}个, 含"nan"文本: {nan_count}个')

print()
print('=== 前10行数据 ===')
headers = [cell.value for cell in ws[1]]
print(f"{'学号':<10}{'姓名':<8}{'联系电话':<15}{'填写日期':<12}{'省份':<8}")
for row in range(2, 12):
    vals = [
        ws[f'A{row}'].value or '',
        ws[f'B{row}'].value or '',
        ws[f'G{row}'].value or '',
        ws[f'F{row}'].value or '',
        ws[f'E{row}'].value or '',
    ]
    print(f"{str(vals[0]):<10}{str(vals[1]):<8}{str(vals[2]):<15}{str(vals[3]):<12}{str(vals[4]):<8}")

print()
if all_ok:
    print('✅ 验证通过！所有空值在Excel中是空单元格，没有变成"nan"文本')
else:
    print('❌ 验证失败！存在"nan"文本')
